"""Extração de texto por página de PDF, DOCX e TXT/MD."""
import os
import re
from typing import List


def _clean(t: str) -> str:
    t = t.replace("\x00", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def extract_pages(path: str) -> List[str]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext in (".txt", ".md", ".html", ".htm"):
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            raw = fh.read()
        if ext in (".html", ".htm"):
            raw = re.sub(r"<[^>]+>", " ", raw)
        return _paginate_text(raw)
    return []


def _pdf(path: str) -> List[str]:
    from pypdf import PdfReader
    pages = []
    reader = PdfReader(path)
    for p in reader.pages:
        try:
            pages.append(_clean(p.extract_text() or ""))
        except Exception:
            pages.append("")
    return pages


def _docx(path: str) -> List[str]:
    import docx
    d = docx.Document(path)
    parts = []
    for para in d.paragraphs:
        parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return _paginate_text("\n".join(parts))


def _paginate_text(raw: str, max_chars: int = 3500) -> List[str]:
    """Documentos sem paginação física: quebra por marcador de página ou por tamanho."""
    if "\f" in raw:
        return [_clean(x) for x in raw.split("\f") if x.strip()]
    paras = [p for p in re.split(r"\n\s*\n", raw)]
    pages, cur = [], ""
    for p in paras:
        if len(cur) + len(p) > max_chars and cur:
            pages.append(_clean(cur))
            cur = ""
        cur += p + "\n\n"
    if cur.strip():
        pages.append(_clean(cur))
    return pages or [""]
